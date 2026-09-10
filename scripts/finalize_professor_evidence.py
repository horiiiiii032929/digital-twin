"""Validate and document the professor evidence archive; no model calls."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sqlite3
import zipfile
from pathlib import Path
from urllib.parse import quote, unquote
import xml.etree.ElementTree as ET

from build_professor_evidence import (
    ROOT,
    DEST,
    BUILD,
    DECK,
    GROUPS,
    digest,
    aid,
    jdump,
    SECRET,
    ingest_json,
)

QUERIES = [
    (
        "01_evidence_index",
        "Evidence groups and their slide and report locations",
        "SELECT * FROM evidence_groups ORDER BY evidence_id",
    ),
    (
        "02_file_catalog",
        "Every catalogued source, dataset, record and output",
        "SELECT * FROM files ORDER BY original_path",
    ),
    (
        "03_run_catalog",
        "Recorded run identities and decisions; distinct artifacts can describe the same run",
        "SELECT r.*,f.package_path FROM runs r JOIN files f USING(file_id) ORDER BY run_id,file_id",
    ),
    (
        "04_gaps",
        "Unresolved paths and intentionally excluded artifacts",
        "SELECT original_path,status,note FROM unresolved_files ORDER BY status,original_path",
    ),
    (
        "05_factual_responses",
        "Actual saved responses from the five fresh factual-comparison arms",
        "SELECT original_path,json_extract(payload_json,'$.case_id') AS case_id,json_extract(payload_json,'$.action') AS action,json_extract(payload_json,'$.answer') AS answer FROM json_records WHERE original_path LIKE 'reports/generated/final-cross-method-factual-confirmation-001/%/responses.sqlite3' AND locator LIKE 'table:responses/payload:%' ORDER BY original_path,case_id",
    ),
    (
        "06_prediction_results",
        "Recalculate history-weighted Brier means from 480 saved open-loop histories",
        "SELECT count(*) AS histories,sum(json_extract(payload_json,'$.observations')) AS observations,avg(json_extract(payload_json,'$.brier.count')) AS count_brier,avg(json_extract(payload_json,'$.brier.decay')) AS decay_brier,avg(json_extract(payload_json,'$.brier.bkt')) AS bkt_brier,avg(json_extract(payload_json,'$.brier.pfa')) AS pfa_brier FROM json_records WHERE original_path='reports/generated/post-report-learner-policy-001-local-001/open-loop-histories.jsonl'",
    ),
    (
        "07_intervention_results",
        "Recalculate all 17 conditions from their saved histories",
        "SELECT json_extract(payload_json,'$.condition') AS condition,count(*) AS histories,avg(json_extract(payload_json,'$.final_hidden_mastery')) AS final_hidden_mastery,avg(json_extract(payload_json,'$.messages_sent')) AS messages_sent,sum(json_extract(payload_json,'$.eligibility_violations')) AS eligibility_violations FROM json_records WHERE original_path='reports/generated/post-report-learner-policy-001-local-001/histories.jsonl' GROUP BY condition ORDER BY condition",
    ),
    (
        "08_reviewer_ratings",
        "All 160 reviewer ratings including failed validation outcomes",
        "SELECT json_extract(payload_json,'$.id') AS output_id,json_extract(payload_json,'$.phase') AS phase,json_extract(payload_json,'$.arm') AS arm,json_extract(payload_json,'$.repetition') AS repetition,json_extract(payload_json,'$.status') AS status,json_extract(payload_json,'$.review.overall') AS overall,payload_json FROM json_records WHERE original_path='research/05_evaluation/records/post-report-final-selection-004-mini-review-live-001.json' AND locator LIKE '$.per_rating[%]' ORDER BY phase,output_id,repetition",
    ),
    (
        "09_diagnostic_acceptance",
        "Recalculate acceptable-in-both counts; failed reviewer gate still applies",
        "WITH ratings AS (SELECT json_extract(payload_json,'$.id') AS output_id,json_extract(payload_json,'$.arm') AS arm,count(*) AS repeats,sum(CASE WHEN json_extract(payload_json,'$.review.overall')='acceptable' THEN 1 ELSE 0 END) AS acceptable FROM json_records WHERE original_path='research/05_evaluation/records/post-report-final-selection-004-mini-review-live-001.json' AND locator LIKE '$.per_rating[%]' AND json_extract(payload_json,'$.phase')='product' GROUP BY output_id,arm) SELECT arm,count(*) AS outputs,sum(CASE WHEN repeats=2 AND acceptable=2 THEN 1 ELSE 0 END) AS acceptable_twice FROM ratings GROUP BY arm ORDER BY arm",
    ),
    (
        "10_sqlite_exclusions",
        "Authentication/secret exclusions by original database",
        "SELECT f.original_path,e.locator,e.reason,e.row_count FROM exclusions e JOIN files f USING(file_id) ORDER BY f.original_path,e.locator",
    ),
    (
        "11_text_search",
        "Example full-text search; replace the quoted term",
        "SELECT f.file_id,f.original_path,substr(t.text,max(1,instr(lower(t.text),'brier')-80),300) AS excerpt FROM texts t JOIN files f USING(file_id) WHERE lower(t.text) LIKE '%brier%' ORDER BY f.original_path",
    ),
    (
        "12_json_search",
        "Example query into a material PDF page or record; change path/term as needed",
        "SELECT original_path,locator,json_extract(payload_json,'$.text') AS page_text FROM json_records WHERE original_path LIKE 'academia_vault/%' AND lower(json_extract(payload_json,'$.text')) LIKE '%data access%' ORDER BY original_path,locator",
    ),
]


def csv_query(c, p, sql, parameters=()):
    p.parent.mkdir(parents=True, exist_ok=True)
    cursor = c.execute(sql, parameters)
    with p.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(x[0] for x in cursor.description)
        w.writerows(cursor)


def checks_and_snapshots(c):
    for name, description, sql in QUERIES:
        csv_query(c, DEST / "snapshots" / f"{name}.csv", sql)
    (DEST / "database/queries.sql").write_text(
        "\n\n".join(f"-- {name}: {desc}\n{sql};" for name, desc, sql in QUERIES) + "\n",
        encoding="utf-8",
    )

    def check(id, passed, description, actual):
        c.execute(
            "INSERT OR REPLACE INTO checks VALUES (?,?,?,?)",
            (id, "PASS" if passed else "UNRESOLVED", description, jdump(actual)),
        )

    row = c.execute(QUERIES[5][2]).fetchone()
    expected = (480, 7535, 0.25508, 0.24821, 0.25141, 0.26588)
    passed = row[:2] == expected[:2] and all(
        abs(a - b) < 0.0000051 for a, b in zip(row[2:], expected[2:])
    )
    check(
        "saved-open-loop-brier",
        passed,
        "480 histories, 7535 observations; four Brier means match the slide/report record to five decimals",
        row,
    )
    interventions = c.execute(QUERIES[6][2]).fetchall()
    check(
        "saved-intervention-grid",
        len(interventions) == 17 and sum(r[1] for r in interventions) == 8160,
        "17 conditions and 8160 stored intervention histories",
        [r[:2] for r in interventions],
    )
    by = {r[0]: r for r in interventions}
    if all(x in by for x in ["count+conditional", "decay+conditional", "bkt+value"]):
        base = by["count+conditional"]
        d = by["decay+conditional"]
        b = by["bkt+value"]
        values = {
            "decay_mastery_delta": d[2] - base[2],
            "bkt_value_mastery_delta": b[2] - base[2],
            "bkt_value_message_delta": b[3] - base[3],
        }
        check(
            "saved-intervention-differences",
            abs(values["decay_mastery_delta"] - 0.00112) < 0.0000051
            and abs(values["bkt_value_mastery_delta"] - 0.02885) < 0.0000051,
            "Saved means match recorded simulator mastery differences to five decimals",
            values,
        )
    ratings = c.execute(QUERIES[8][2]).fetchall()
    check(
        "saved-reviewer-diagnostics",
        ratings == [("v19-luna-luna-medium", 24, 17), ("v4", 24, 13)],
        "Recorded acceptable-in-both counts; this does not change the failed reviewer gate",
        ratings,
    )
    factual = c.execute(
        "SELECT original_path,count(*) FROM json_records WHERE original_path LIKE 'reports/generated/final-cross-method-factual-confirmation-001/%/responses.sqlite3' AND locator LIKE 'table:responses/payload:%' GROUP BY original_path"
    ).fetchall()
    check(
        "factual-output-counts",
        len(factual) == 5 and all(n == 1000 for _, n in factual),
        "Five saved response ledgers, 1000 responses per arm; count verification only",
        factual,
    )
    # Preserve all other reported values without an unsupported claim of re-scoring.
    check(
        "historical-disputed-score",
        False,
        "Whole-corpus success percentage remains disputed; original ledger missing in earlier audit; no substitute score computed",
        {"source": "docs/post-report-historical-metric-discrepancy-2026-09-08.md"},
    )
    c.commit()
    csv_query(
        c, DEST / "snapshots/verification.csv", "SELECT * FROM checks ORDER BY check_id"
    )
    for (g,) in c.execute(
        "SELECT evidence_id FROM evidence_groups ORDER BY evidence_id"
    ).fetchall():
        csv_query(
            c,
            DEST / "snapshots" / f"{g}_materials.csv",
            "SELECT file_id,original_path,package_path,status,sha256 FROM group_inventory WHERE evidence_id=? AND (category IN ('materials','datasets') OR original_path LIKE '%/instruments/%') ORDER BY original_path",
            (g,),
        )
        csv_query(
            c,
            DEST / "snapshots" / f"{g}_records.csv",
            "SELECT file_id,original_path,package_path,status,note FROM group_inventory WHERE evidence_id=? ORDER BY original_path",
            (g,),
        )
    (DEST / "database/data-dictionary.md").write_text("""# Evidence database dictionary

This is SQLite SQL, not PostgreSQL. Open evidence.sqlite in DB Browser for SQLite.
Every table is a new index over historical files; it does not replace a recorded evaluation.

| Table or view | Row grain |
|---|---|
| evidence_groups | One linked study or implementation/document group |
| files | One original file path, with original and delivered hashes |
| group_files | One group-to-file link; related files may belong to multiple groups |
| file_links | One explicitly recorded path reference |
| texts | One text document or source-code file |
| payloads | One distinct JSON value, deduplicated by SHA-256 |
| records | One JSON root, array member, JSONL line, CSV row, PDF page or SQLite row |
| json_records | Joins records, payloads and files into a queryable JSON view |
| case_records | Records with an explicit top-level case_id; not every dataset uses this schema |
| response_records | Case records with an answer or response field |
| runs | One source record with a run_id; the same run may have several source records |
| exclusions | One excluded authentication table or secret-containing row |
| checks | One archive verification or saved-data arithmetic check |
| document_locations | One presentation slide or submitted source section/file |
| unresolved_files | Missing, excluded or otherwise unavailable paths |

Do not count json_records indiscriminately: JSON roots and indexed array members overlap.
Filter by original_path and locator before aggregating. JSONL locators are physical line
numbers. SQLite row locators identify traversal positions in the export, not original
primary keys; primary keys remain in payload_json. Binary SQLite columns are represented
by hashes and byte counts, with exact bytes retained in the corresponding .sqlite3.gz
companion. Decompress that companion to query its binary values directly.

Use json_extract(payload_json, '$.field'), json_each for arrays, and json_tree for nested
content. Missing JSON fields are NULL, not zero. Native nested *_json columns have been
parsed as JSON objects in this index; companion SQLite exports preserve original values.

Sanitized SQLite companions retain table names, columns, values and BLOBs. Authentication
tables and detected secret rows are excluded. Indexes, triggers and constraints are not
replayed. They are research data exports, not deployable application backups.

Paths are relative to the extracted archive root. Original absolute paths are provenance
only. Current code is labelled as current context; historical revision and dirty state
remain in the original run record. Hashes establish identity, not scientific validity.
""")


def hyperlink(paragraph, label, target):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    h = OxmlElement("w:hyperlink")
    h.set(
        qn("r:id"),
        paragraph.part.relate_to(
            quote(target, safe="/:#?=&%"), RT.HYPERLINK, is_external=True
        ),
    )
    r = OxmlElement("w:r")
    rp = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rp.append(color)
    r.append(rp)
    t = OxmlElement("w:t")
    t.text = label
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def docx_index(c):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(11.69)
    sec.page_height = Inches(8.27)
    sec.top_margin = sec.bottom_margin = Inches(0.55)
    sec.left_margin = sec.right_margin = Inches(0.55)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        st = doc.styles[name]
        st.font.name = "Arial"
        st.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(9)
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)
    doc.styles["Title"].font.size = Pt(20)
    doc.add_paragraph("Digital Twin Evidence Index", "Title")
    doc.add_paragraph("Hikaru Rawin Horinouchi | Reference guide for professor review")
    doc.add_paragraph(
        "This index connects the final 70-slide presentation dated 9 September 2026 and the report submitted on 6 September 2026 to their datasets, source materials and saved evaluation results. Extract the complete ZIP before using the relative file links. The submitted report and later presentation use different experiments where indicated."
    )
    p = doc.add_paragraph()
    hyperlink(p, "Open SQLite database", "database/evidence.sqlite")
    p.add_run("  |  ")
    hyperlink(p, "SQL dump", "database/evidence.sql")
    p.add_run("  |  ")
    hyperlink(p, "Query list", "database/queries.sql")
    p.add_run("  |  ")
    hyperlink(p, "Complete file catalog", "snapshots/02_file_catalog.csv")
    p.add_run("  |  ")
    hyperlink(p, "Gaps and exclusions", "snapshots/04_gaps.csv")
    doc.add_paragraph(
        "The database includes searchable text and structured records; original materials are linked from the catalogs. SQLite companions ending in .gz must be decompressed before opening. Recorded results remain historical evidence: missing ledgers, failed gates and disputed values are not repaired by this archive."
    )
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.63, 2.02, 0.75, 2.13, 2.25, 2.79]
    widths = [x * 10.59 / sum(widths) for x in widths]
    headers = [
        "ID",
        "Evaluation or experiment",
        "Slide(s)",
        "Submitted report location",
        "Dataset and source materials",
        "Result and actual outputs",
    ]
    for col, w in zip(table.columns, widths):
        col.width = Inches(w)
    for cell, text, w in zip(table.rows[0].cells, headers, widths):
        cell.text = text
        cell.width = Inches(w)
    trpr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    trpr.append(repeat)
    groups = c.execute(
        "SELECT * FROM evidence_groups ORDER BY CASE WHEN evidence_id LIKE 'E%' THEN 0 ELSE 1 END,evidence_id"
    ).fetchall()
    for gid, title, slides, loc in groups:
        cells = table.add_row().cells
        for cell, w in zip(cells, widths):
            cell.width = Inches(w)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[0].text = gid
        cells[1].text = title
        cells[2].text = (
            "See full slide list in index CSV" if gid == "D02" else slides or "—"
        )
        # The full location list is also in the machine-readable evidence index.
        locs = loc.split("; ")
        compact = [x for x in locs if x.startswith("Appendix Evidence Index")]
        other = [x for x in locs if not x.startswith("Appendix Evidence Index")]
        cells[3].text = "\n".join(other[:1] + compact) if compact else loc
        hyperlink(
            cells[4].paragraphs[0],
            "Materials and dataset catalog",
            f"snapshots/{gid}_materials.csv",
        )
        cells[4].add_paragraph("Includes file locations, availability and hashes.")
        p = cells[5].paragraphs[0]
        if gid.startswith("E"):
            run = GROUPS[int(gid[1:]) - 1][2][0]
            found = c.execute(
                "SELECT package_path FROM files WHERE original_path=?",
                (f"research/05_evaluation/{run}-results.md",),
            ).fetchone()
            if found and found[0]:
                hyperlink(p, "Result summary", found[0])
                p.add_run(" | ")
        hyperlink(p, "All records and outputs", f"snapshots/{gid}_records.csv")
        unresolved = c.execute(
            "SELECT count(*) FROM group_inventory WHERE evidence_id=? AND status='unresolved_reference'",
            (gid,),
        ).fetchone()[0]
        if unresolved:
            cells[5].add_paragraph(
                f"{unresolved} unresolved referenced path(s); see catalog. Counts include supporting historical references."
            )
        if gid == "E25":
            cells[5].add_paragraph(
                "Historical score disputed. Do not cite it as verified performance."
            )
        if gid == "E11":
            cells[5].add_paragraph(
                "Reviewer gate failed; quality ratings remain diagnostic."
            )
        if gid in {"E13", "E21"}:
            cells[5].add_paragraph(
                "Synthetic simulation; not measured student learning."
            )
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            pr = cell._tc.get_or_add_tcPr()
            b = OxmlElement("w:tcBorders")
            for name in ["top", "left", "bottom", "right"]:
                e = OxmlElement("w:" + name)
                e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), "4")
                e.set(qn("w:color"), "D9D9D9")
                b.append(e)
            pr.append(b)
            m = OxmlElement("w:tcMar")
            for name in ["top", "left", "bottom", "right"]:
                e = OxmlElement("w:" + name)
                e.set(qn("w:w"), "80")
                e.set(qn("w:type"), "dxa")
                m.append(e)
            pr.append(m)
            if ri == 0:
                sh = OxmlElement("w:shd")
                sh.set(qn("w:fill"), "E7E7E7")
                pr.append(sh)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
        if ri:
            pr = row._tr.get_or_add_trPr()
            pr.append(OxmlElement("w:cantSplit"))
    doc.add_paragraph("Queries and verification", "Heading 1")
    p = doc.add_paragraph(
        "Open the SQLite file in DB Browser for SQLite and use Execute SQL. "
    )
    hyperlink(p, "Setup and restore instructions", "README.txt")
    p.add_run(" explain how to recreate the database from the SQL dump.")
    p = doc.add_paragraph()
    hyperlink(p, "Saved query results", "snapshots/01_evidence_index.csv")
    p.add_run(" and ")
    hyperlink(p, "verification results", "snapshots/verification.csv")
    p.add_run(
        " were generated from this database. The query list includes inputs and responses, reviewer ratings, learner prediction means, intervention comparisons, source-text search and exclusions."
    )
    p = sec.footer.paragraphs[0]
    p.alignment = 2
    p.add_run("Evidence index | ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)
    doc.core_properties.title = "Digital Twin Evidence Index"
    doc.core_properties.author = "Hikaru Rawin Horinouchi"
    # The bundled default Word template may carry a blue Title border.
    for root in [doc.styles.element, doc.element]:
        for border in list(root.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    doc.save(DEST / "Evidence Index.docx")


def finalize():
    c = sqlite3.connect(DEST / "database/evidence.sqlite")
    # Add the directly cited report trace, which is outside its A--L study index.
    c.execute(
        "INSERT OR REPLACE INTO evidence_groups VALUES (?,?,?,?)",
        (
            "E27",
            "Report historical multi concept trace",
            "",
            "Appendix Historical multi concept autonomy trace",
        ),
    )
    c.execute(
        "INSERT OR IGNORE INTO group_files SELECT 'E27',file_id FROM files WHERE original_path LIKE '%governed-full-autonomy-v2-1-multi-concept-confirmation-025%'"
    )
    # Retain a metric record that initially matched a token substring within 'risk-'.
    import shutil

    for fid, k in c.execute(
        "SELECT file_id,original_path FROM files WHERE status='excluded_sensitive'"
    ).fetchall():
        source = ROOT / k
        if (
            source.is_file()
            and source.suffix == ".json"
            and not SECRET.search(source.read_text())
        ):
            rel = "records/" + k
            out = DEST / rel
            out.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, out)
            obj = json.loads(source.read_text())
            ingest_json(c, fid, obj)
            c.execute(
                "UPDATE files SET package_path=?,status=?,packaged_sha256=?,note=? WHERE file_id=?",
                (
                    rel,
                    "included",
                    digest(out),
                    "Original JSON; secret scan false positive resolved at token boundary",
                    fid,
                ),
            )
    # Propagate explicit file references so shared prior discovery cannot hide a gap
    # from another study that cites the same source record.
    c.execute("""WITH RECURSIVE closure(evidence_id,file_id) AS (
 SELECT evidence_id,file_id FROM group_files UNION
 SELECT closure.evidence_id,file_links.target_file_id FROM closure JOIN file_links ON file_links.source_file_id=closure.file_id
 ) INSERT OR IGNORE INTO group_files SELECT evidence_id,file_id FROM closure""")
    c.commit()
    checks_and_snapshots(c)
    (DEST / "README.txt").write_text(
        """DIGITAL TWIN PROFESSOR EVIDENCE ARCHIVE

1. Extract the complete ZIP, keeping its folder structure.
2. Open Evidence Index.docx. Its links resolve within the extracted folder.
3. Open database/evidence.sqlite with DB Browser for SQLite (sqlitebrowser.org).
4. Use Execute SQL and open database/queries.sql. Matching results are in snapshots/.

Recreate the database from its SQLite dump using the sqlite3 command-line tool:
  sqlite3 restored.sqlite < database/evidence.sql
Run from the archive root. This is SQLite SQL, not PostgreSQL/pgAdmin SQL.
Alternatively use DB Browser for SQLite: File > Import > Database from SQL file.

Original materials and saved records are under materials/, datasets/, records/ and
outputs/. Use snapshots/02_file_catalog.csv or query files for their exact locations.
The original path is provenance; package_path is the usable relative location.
Files with .sqlite3.gz or .sqlite.gz are sanitized SQLite companions: decompress
with gzip -dk FILE.gz, then open the resulting database. Exact binary checkpoint
values are available there; their hashes and sizes are indexed in evidence.sqlite.

Read database/data-dictionary.md before aggregating records. JSON roots and array
members overlap. A row count is not automatically a count of independent cases.
The archive contains related historical controls, failures and corrections as well
as primary results. Group membership records discovery links, not a claim that
every supporting record was displayed on a slide.

The original experiments were not rerun. Recorded metrics are preserved. The
verification CSV identifies the limited arithmetic checks reproduced from saved
rows. Reproduction of a table does not qualify a failed gate or validate a reviewer.
Missing historical ledgers remain gaps; this is not a fully recoverable archive.

Private professor research review only. Approved course material is not licensed
for public redistribution. Public source licenses and attribution remain with the
source files. Authentication/session material is excluded and recorded separately.
Do not use these research exports as production application backups.

The separately downloaded DOCX requires the extracted archive for local links.
The online download links expire; the extracted archive remains usable offline.
""",
        encoding="utf-8",
    )
    docx_index(c)
    c.commit()
    c.close()
    print("Index and query snapshots created")


def dump_verify():
    c = sqlite3.connect(DEST / "database/evidence.sqlite")
    assert c.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    assert not c.execute("PRAGMA foreign_key_check").fetchall()
    with (DEST / "database/evidence.sql").open("w", encoding="utf-8") as f:
        for line in c.iterdump():
            f.write(line + "\n")
    rebuilt = BUILD / "restored.sqlite"
    if rebuilt.exists():
        rebuilt.unlink()
    r = sqlite3.connect(rebuilt)
    r.execute("PRAGMA journal_mode=MEMORY")
    r.execute("PRAGMA synchronous=OFF")
    with (DEST / "database/evidence.sql").open(encoding="utf-8") as f:
        buffer = ""
        for line in f:
            buffer += line
            if sqlite3.complete_statement(buffer):
                r.execute(buffer)
                buffer = ""
    r.commit()
    assert r.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    tables = [
        x[0]
        for x in c.execute(
            "SELECT name FROM sqlite_master WHERE type='table' ORDER BY name"
        )
    ]
    comparison = {}
    for table in tables:

        def tablehash(db):
            h = hashlib.sha256()
            count = 0
            for row in db.execute('SELECT * FROM "' + table + '" ORDER BY rowid'):
                h.update(jdump(row).encode())
                h.update(b"\n")
                count += 1
            return count, h.hexdigest()

        a, b = tablehash(c), tablehash(r)
        assert a == b, (table, a, b)
        comparison[table] = {"rows": a[0], "sha256": a[1]}
    for name, _, sql in QUERIES:
        csv_query(r, BUILD / (name + ".csv"), sql)
        assert digest(BUILD / (name + ".csv")) == digest(
            DEST / "snapshots" / (name + ".csv")
        ), name
    with zipfile.ZipFile(DEST / "Evidence Index.docx") as z:
        rels = ET.fromstring(z.read("word/_rels/document.xml.rels"))
        links = []
        for rel in rels:
            if rel.attrib.get("Type", "").endswith("/hyperlink"):
                target = unquote(rel.attrib["Target"])
                if not target.startswith(("http:", "https:", "#")):
                    assert (DEST / target).exists(), target
                    links.append(target)
    # Relocate the actual index and all linked targets using hard links to save disk.
    relocation = BUILD / "relocation-test"
    relocation.mkdir(exist_ok=True)
    import os

    for target in links + ["Evidence Index.docx"]:
        p = relocation / target
        p.parent.mkdir(parents=True, exist_ok=True)
        if not p.exists():
            os.link(DEST / target, p)
        assert digest(p) == digest(DEST / target)
    r.close()
    rebuilt.unlink()
    c.close()
    (BUILD / "restore-validation.json").write_text(
        json.dumps(
            {
                "restored_tables": comparison,
                "query_snapshots_match": len(QUERIES),
                "relative_links_checked": len(links),
                "relocation_check": "passed",
            },
            indent=2,
        )
    )
    print("Dump restore, table hashes, snapshots and relative links verified")


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--dump-verify", action="store_true")
    args = p.parse_args()
    if args.dump_verify:
        dump_verify()
    else:
        finalize()
